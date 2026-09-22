from docx import Document

FILES = [
    r"E:\88888\第二轮作业-王英琦.docx",
    r"E:\88888\AI应用实习生申请-王英琦-山东交通学院-AI应用.docx",
]

for file_path in FILES:
    print(f"\n=====FILE {file_path} =====")
    document = Document(file_path)
    print(
        "PARAGRAPHS",
        len(document.paragraphs),
        "TABLES",
        len(document.tables),
        "SECTIONS",
        len(document.sections),
    )
    for index, paragraph in enumerate(document.paragraphs):
        text = paragraph.text.strip()
        if text:
            print(f"P{index}: {text}")
    for table_index, table in enumerate(document.tables):
        print(f"---TABLE {table_index} {len(table.rows)}x{len(table.columns)}---")
        for row_index, row in enumerate(table.rows):
            values = [cell.text.replace("\n", " / ").strip() for cell in row.cells]
            print(f"R{row_index}: " + " || ".join(values))
