from pathlib import Path
from zipfile import ZipFile
import sys

from docx import Document
from pypdf import PdfReader
import pypdfium2 as pdfium

sys.stdout.reconfigure(encoding="utf-8")
ROOT = Path(r"D:\VentureLab\tmp\review\zhou-shihao-20260922")
ROOT.mkdir(parents=True, exist_ok=True)

docx_path = Path(r"E:\88888\二轮面试回答.docx")
document = Document(docx_path)
print(f"DOCX paragraphs={len(document.paragraphs)} tables={len(document.tables)} sections={len(document.sections)}")
for index, paragraph in enumerate(document.paragraphs):
    text = paragraph.text.strip()
    if text:
        print(f"P{index}: {text}")
for table_index, table in enumerate(document.tables):
    print(f"---TABLE {table_index} {len(table.rows)}x{len(table.columns)}---")
    for row_index, row in enumerate(table.rows):
        values = [cell.text.replace("\n", " / ").strip() for cell in row.cells]
        print(f"R{row_index}: " + " || ".join(values))

with ZipFile(docx_path) as archive:
    media_out = ROOT / "docx-media"
    media_out.mkdir(exist_ok=True)
    media_files = [name for name in archive.namelist() if name.startswith("word/media/")]
    xml = archive.read("word/document.xml")
    print(
        f"DOCX media={len(media_files)} comments={'word/comments.xml' in archive.namelist()} "
        f"insertions={xml.count(b'<w:ins')} deletions={xml.count(b'<w:del')}"
    )
    for name in media_files:
        target = media_out / Path(name).name
        target.write_bytes(archive.read(name))
        print(f"MEDIA {target.name} {target.stat().st_size}")

for label, pdf_path in [
    ("resume", Path(r"E:\88888\周世豪-简历.pdf")),
    ("application", Path(r"E:\88888\AI应用实习生申请.pdf")),
]:
    reader = PdfReader(str(pdf_path))
    print(f"\n=====PDF {label} pages={len(reader.pages)}=====")
    for index, page in enumerate(reader.pages, start=1):
        print(f"\n--- {label.upper()} PAGE {index} ---")
        print(page.extract_text() or "")
    pdf = pdfium.PdfDocument(str(pdf_path))
    out_dir = ROOT / f"{label}-pages"
    out_dir.mkdir(exist_ok=True)
    for index in range(len(pdf)):
        image = pdf[index].render(scale=1.8).to_pil()
        image.save(out_dir / f"page-{index + 1}.png")
