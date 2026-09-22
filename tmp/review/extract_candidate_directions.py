from pathlib import Path
import re

from pypdf import PdfReader
from docx import Document


FILES = [
    Path(r"E:\88888\黄威-湖南工商大学-AI应用实习申请问卷.pdf.pdf"),
    Path(r"E:\88888\第二轮作业＋黄威 (1).pdf"),
    Path(r"E:\88888\苏琼-湖南工学院-简历.pdf"),
    Path(r"E:\88888\第二轮作业-苏琼 (1).pdf"),
    Path(r"E:\88888\AI应用实习生申请-王英琦-山东交通学院-AI应用.docx"),
    Path(r"E:\88888\第二轮作业-王英琦.docx"),
    Path(r"E:\88888\谭北川_湖南涉外经济学院_简历.pdf"),
    Path(r"E:\88888\第二轮作业-张尔得.pdf"),
    Path(r"E:\88888\张尔得-湖南师范大学-简历.pdf"),
    Path(r"E:\88888\AI应用实习生申请.pdf"),
    Path(r"E:\88888\二轮面试回答.docx"),
    Path(r"E:\88888\周世豪-简历.pdf"),
    Path(r"E:\88888\第二轮作业＋邹振豪.pdf"),
    Path(r"E:\88888\邹振豪-27年应届生-大模型应用开发AI Agent工程师求职简历 (2) (1).pdf"),
]


def read_file(path: Path) -> str:
    if path.suffix.lower() == ".pdf":
        return "\n".join((page.extract_text() or "") for page in PdfReader(path).pages)
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


PATTERNS = [
    r"目前所在城市.{0,250}",
    r"当前所在城市.{0,250}",
    r"学校.{0,220}",
    r"方向选择.{0,800}",
    r"优先方向.{0,500}",
    r"方向一.{0,500}",
    r"方向二.{0,500}",
    r"更希望.{0,500}",
    r"感兴趣.{0,500}",
    r"三个月后.{0,500}",
    r"到岗.{0,350}",
    r"最晚.{0,250}",
    r"共同条件.{0,800}",
    r"六级.{0,220}",
    r"CET.{0,220}",
    r"英语.{0,220}",
]

for path in FILES:
    if not path.exists():
        print(f"\n### MISSING {path.name}")
        continue
    text = re.sub(r"\s+", " ", read_file(path))
    print(f"\n### {path.name}")
    seen = set()
    for pattern in PATTERNS:
        for match in re.finditer(pattern, text, flags=re.I):
            excerpt = match.group(0)[:900]
            if excerpt not in seen:
                print(excerpt)
                seen.add(excerpt)
